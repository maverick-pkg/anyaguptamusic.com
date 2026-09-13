#!/usr/bin/env python3
"""Generate the September 18 checklist (night-before version) as a Word document.

Writes  ~/Dropbox (Personal)/Music/Anya/YouTube/SEPT 18 SWAP/SEPTEMBER 18 CHANGEOVER - checklist.docx
and first moves the 7 Sept version to  SEPT 18 SWAP/Old versions/  (once).
Regenerate from this script; don't hand-edit the .docx.
Run:  /usr/bin/python3 tools/release/make_checklist_docx.py
"""
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SWAP = Path.home() / "Dropbox (Personal)/Music/Anya/YouTube/SEPT 18 SWAP"
OUT = SWAP / "SEPTEMBER 18 CHANGEOVER - checklist.docx"
OLD = SWAP / "Old versions" / "SEPTEMBER 18 CHANGEOVER - checklist (7 Sep version).docx"
INK, SOFT, PEACH = RGBColor(0x21, 0x1D, 0x20), RGBColor(0x5A, 0x52, 0x57), RGBColor(0xB8, 0x5A, 0x28)

DONE = [
    "Spotify editorial pitch for “Not For You” — sent Aug 24.",
    "All 7 Spotify Canvases uploaded (you confirmed, Sept 13).",
    "Writer credits confirmed; Spotify bio updated; YouTube files built.",
    "YouTube descriptions for the four new songs — written (in this folder, “description - <song> - OUT NOW.txt”).",
    "Website release version built by Claude and independently reviewed (kept unpublished until the album is streaming).",
]

BEFORE = [
    ("You", "Fix two credit lines that are live on YouTube right now (she co-wrote her songs):",
     ["Channel description: change “7 songs i wrote” to “7 songs i co-wrote”. Nothing else.",
      "You’re Original video: paste the corrected “description - You’re Original - REVISED.txt” (in Music/Anya/YouTube/). It no longer says “Written and performed by Anya Gupta” as if she wrote it alone."]),
    ("You", "YouTube — upload the four held lyric videos as PRIVATE. Don’t publish yet.",
     ["Files: Music/Anya/YouTube/Lyric videos - Official/ — Not For You, Let You Be, Breakup with My Ego, Life Vest.",
      "Title exactly: Anya Gupta - <Song> (Official Lyric Video)   (Claude finds the video by this title).",
      "Description: paste the whole “description - <Song> - OUT NOW.txt” file from this folder.",
      "Thumbnail: “thumbnail - <Song>.jpg” and subtitles “<Song>.srt” — both in Music/Anya/YouTube/.",
      "Add to the playlist “Anya Gupta – Becoming | Official Videos”; end screen on the outro card; a card at ~25s.",
      "Leave every video on Private."]),
    ("You", "YouTube — set the playlist order: You’re Original, Not For You, Tat Twam Asi, Let You Be, Breakup with My Ego, Awaken the Lights, Life Vest.", []),
    ("You", "Spotify Artist Pick — pin “You’re Original” now, with a message such as “my debut album Becoming is out sept 18 ✦”.",
     ["Spotify only lets you pin music that is already out, so the album itself goes up Thursday night.",
      "artists.spotify.com → View Profile → + under Artist Pick → search the song → add the message → Save."]),
    ("You", "Mailchimp — send “Becoming is out Friday — follow Anya on Spotify so it shows up for you.” Draft the Friday “out now” email now too.", []),
    ("You", "Tell Claude if WREG gives you the Live at 9 airdate — the site line and listing get the real date.", []),
]

THURSDAY = [
    ("You", "Morning: the “tomorrow” post (video 4-3).", []),
    ("You", "7:00 PM — open Spotify and Apple Music and search Anya Gupta. Is Becoming there and playing?",
     ["DistroKid’s countdown ends at 7:00 PM Memphis time. If it isn’t there, check again at 11:00 PM.",
      "Don’t start the rest until you can actually play it."]),
    ("You", "YouTube — switch the four lyric videos from Private to Public. Do this BEFORE starting Claude.",
     ["The site shows the Not For You lyric video, and Claude can only find it once it is public."]),
    ("You + Claude", "Start a Claude session in the anya-site folder and say: “Becoming is live — run the release.”",
     ["Claude pulls the Spotify, Apple and YouTube links itself, fills them in, runs the checks, and shows you the pages.",
      "If Apple Music is slow to list the album, Claude may ask you for the Apple links (Music app → Share → Copy Link).",
      "On your “publish”, the site goes live: album page, four new lyrics pages, “out now” homepage, press kit, pre-save and thank-you pages.",
      "About 30–45 minutes. Claude also refreshes the press-kit PDF in Music/Anya/Biography/."]),
    ("You", "YouTube — the date swap (files in this folder):",
     ["Channel banner → “channel banner E - OUT NOW.jpg”.",
      "Channel description → “channel bio - OUT NOW.txt”.",
      "Descriptions of You’re Original, Tat Twam Asi and Awaken the Lights → their “OUT NOW” files.",
      "Album trailer description → “description - Album Trailer - OUT NOW.txt”; delete any pinned comment that says pre-save.",
      "Optional: make “Not For You” the featured video."]),
    ("You", "Spotify for Artists — Artist Pick → the album Becoming (message e.g. “my debut album is out now ✦”).", []),
    ("You", "Spotify for Artists bio — change only “is out september 18” to “is out now”. Keep her words and ★ marks.", []),
    ("You", "DistroKid — HyperFollow bio: “out September 18” → “out now”.", []),
    ("You", "KOMI link page — “pre-save becoming” → the album’s listen link (Claude gives it to you); retitle the page.", []),
    ("You", "Optional: a quick “it’s out!” story. The big post is Friday morning.", []),
]

FRIDAY = [
    ("You", "Full-song post (video 10) on Reels, TikTok and Shorts. Captions switch to “listen everywhere”.", []),
    ("You", "Mailchimp — send the “out now” email.", []),
    ("You + Claude", "MusicBrainz — add the album (Claude gives the exact entries) and Google Search Console — request indexing (Claude lists the pages).", []),
    ("You", "YouTube — if a Content ID claim lands on a lyric video or the trailer, clear it at distrokid.com/youtubeAllowlist. Claims can only be cleared after they appear.", []),
]

AFTER = [
    ("You", "Sat 19: video 7 (second full version).  Sun 20: “use the sound” post.", []),
    ("You + Claude", "Tue 22 (her birthday): 15 → 16 — Claude updates the site and press kit; you update the Spotify and HyperFollow bios.", []),
]

NOTES = [
    "Why Thursday night: DistroKid’s pre-save countdown ends Thursday 7:00 PM Memphis time, and Spotify lists the US release date as Thursday, Sept 17. The official date stays September 18 — nothing printed needs to change.",
    "Don’t post a full song before it is streaming.",
    "No stream or view numbers in posts or press without Claude checking where they came from first.",
]


def para(doc, text, size=11, bold=False, color=INK, italic=False, space_after=4, indent=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size, r.bold, r.italic, r.font.color.rgb = Pt(size), bold, italic, color
    p.paragraph_format.space_after = Pt(space_after)
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    return p


def section(doc, title, sub=None):
    p = para(doc, title, size=14, bold=True, color=PEACH, space_after=2)
    p.paragraph_format.space_before = Pt(10)
    if sub:
        para(doc, sub, size=10, color=SOFT, italic=True, space_after=6)


def items(doc, rows):
    for who, text, subs in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2 if subs else 5)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        a = p.add_run("☐  ")
        a.font.size = Pt(12)
        w = p.add_run(f"{who}:  ")
        w.bold, w.font.size, w.font.color.rgb = True, Pt(10), SOFT
        t = p.add_run(text)
        t.font.size = Pt(11)
        for s in subs:
            para(doc, "–  " + s, size=10, color=SOFT, space_after=1, indent=0.55)
        if subs:
            doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def main():
    if OUT.exists() and not OLD.exists():
        OLD.parent.mkdir(exist_ok=True)
        shutil.move(str(OUT), str(OLD))
        print("moved old version ->", OLD.relative_to(SWAP))
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.7)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)

    h = para(doc, "September 18 checklist — Becoming", size=20, bold=True, space_after=0)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para(doc, "Night-before version · most of the work happens Thursday night, once the album is streaming", size=11,
         color=PEACH, italic=True, space_after=6)
    para(doc, "Files mentioned are in Dropbox → Music/Anya/YouTube/SEPT 18 SWAP/ unless another folder is named.",
         size=10, color=SOFT, space_after=4)

    section(doc, "Already done")
    for d in DONE:
        para(doc, "✓  " + d, size=10.5, color=SOFT, space_after=2)
    section(doc, "Before Thursday (Mon 14 – Wed 16)")
    items(doc, BEFORE)
    section(doc, "Thursday, Sept 17 — the night before (the main block)",
            "Start at 7:00 PM, once you can play the album on Spotify.")
    items(doc, THURSDAY)
    section(doc, "Friday, Sept 18 — release day (light)")
    items(doc, FRIDAY)
    section(doc, "After")
    items(doc, AFTER)
    section(doc, "Notes")
    for n in NOTES:
        para(doc, "•  " + n, size=10, color=SOFT, space_after=3)
    doc.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    main()
